"""有范围和数量预算的只读观察，供 Codex 理解布局后编写映射。"""

from pathlib import Path
import openpyxl
from openpyxl.utils.cell import range_boundaries

from workbook_images import WorkbookPackage


def inspect_source(source, sheet=None, cell_range="A1:P8", limit=10, offset=0, max_cells=80, max_chars=100):
    if not 1 <= limit <= 50 or offset < 0 or not 1 <= max_cells <= 200 or not 20 <= max_chars <= 500:
        raise ValueError("Invalid inspection budget")
    source = Path(source).resolve()
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet is None:
            selected = list(book)[offset:offset + limit]
            package = WorkbookPackage(source)
            try:
                return {"file": source.name, "sheet_count": len(book.sheetnames), "offset": offset, "limit": limit,
                        "sheets": [{"name": s.title, "declared_rows": s.max_row, "declared_columns": s.max_column,
                                    **package.sheet_extent(s.title)} for s in selected],
                        "more": offset + limit < len(book.sheetnames)}
            finally:
                package.close()
        if sheet not in book.sheetnames:
            raise ValueError("Unknown sheet; inspect sheet names first")
        left, top, right, bottom = range_boundaries(cell_range)
        if not all((left, top, right, bottom)) or not 1 <= bottom - top + 1 <= 50 or not 1 <= right - left + 1 <= 30:
            raise ValueError("Inspect at most 50 rows by 30 columns at a time")
        package = WorkbookPackage(source)
        try:
            formulas, merges = package.sheet_info(sheet)
            values = []
            for row in book[sheet].iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                for cell in row:
                    if cell.value is None and getattr(cell, "coordinate", None) not in formulas:
                        continue
                    value = cell.value if cell.value is not None else formulas[cell.coordinate]
                    text = str(value)
                    values.append({"cell": cell.coordinate, "value": text[:max_chars], "truncated": len(text) > max_chars})
            relevant = []
            for merged in merges:
                a, b, c, d = range_boundaries(merged)
                if a <= right and c >= left and b <= bottom and d >= top:
                    relevant.append(merged)
            return {"file": source.name, "sheet": sheet, "range": cell_range, "total_nonempty": len(values),
                    "cells": values[offset:offset + max_cells], "more": offset + max_cells < len(values),
                    "merges": relevant[:max_cells], "more_merges": len(relevant) > max_cells}
        finally:
            package.close()
    finally:
        book.close()
