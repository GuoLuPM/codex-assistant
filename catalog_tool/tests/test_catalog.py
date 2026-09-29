import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from catalog_store import Catalog, parse_price


def workbook(path, sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)


class CatalogTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "supplier.xlsx"
        workbook(self.source, {
            "数码": [["产品名称", "供货价", "零售价", "功能特点", "品类"],
                   ["H1 蓝牙耳机", 88, 150, "蓝牙5.3", "数码/耳机"],
                   ["H2 蓝牙耳机", "更新中", 160, "降噪", "数码/耳机"],
                   ["T1 充电线", 10, 20, "附注中出现耳机", "数码/线材"]],
            "家居": [["名称", "供货价", "卖点", "类别"],
                   ["保温杯", 88, "316不锈钢", "家居/杯壶"]],
        })
        self.catalog = Catalog(self.root / "index", {"category_rules": []})

    def tearDown(self):
        self.catalog.close()
        self.temp.cleanup()

    def test_multi_sheet_price_basis_and_compact_search(self):
        self.catalog.index([self.source])
        result = self.catalog.search(query="耳机", price_field="supply_price", minimum=80, maximum=100)
        self.assertEqual([x["name"] for x in result["items"]], ["H1 蓝牙耳机"])
        self.assertNotIn("features", result["items"][0])
        self.assertEqual(result["items"][0]["price"]["label"], "供货价")
        self.assertEqual(self.catalog.search(category="家居")["total"], 1)
        self.assertEqual(self.catalog.facets("category", "数码", limit=1)["total"], 2)
        self.assertEqual(len(self.catalog.facets("category", "数码", limit=1)["items"]), 1)
        self.assertEqual(len(self.catalog.facets("price", "价", limit=1)["items"]), 1)
        self.assertGreater(self.catalog.facets("price", "价", limit=1)["total"], 1)
        with self.assertRaisesRegex(ValueError, "price.field"):
            self.catalog.search(minimum=80)

    def test_features_are_opt_in_and_unknown_price_not_zero(self):
        self.catalog.index([self.source])
        self.assertEqual(self.catalog.search(query="耳机")["total"], 2)
        result = self.catalog.search(query="耳机", scope="all")
        self.assertEqual(result["total"], 3)
        self.assertTrue(any("evidence" in x for x in result["items"]))
        self.assertEqual(self.catalog.search(price_field="supply_price", maximum=0)["total"], 0)
        self.assertIsNone(parse_price("80-150"))
        self.assertIsNone(parse_price("更新中"))
        self.assertEqual(parse_price("￥1,200.50元"), 1200.5)

    def test_incremental_refresh_and_stale_selection(self):
        self.catalog.index([self.source])
        result = self.catalog.search(query="保温杯")
        selected_id = result["items"][0]["id"]
        self.assertEqual(self.catalog.index([self.source])["skipped"], 1)
        workbook(self.source, {"家居": [["名称", "供货价"], ["保温杯", 99]]})
        with self.assertRaisesRegex(ValueError, "changed|stale"):
            self.catalog.stage([selected_id], self.root / "stage")
        self.assertEqual(self.catalog.search()["total"], 0)
        self.catalog.index([self.source])
        refreshed = self.catalog.search(query="保温杯")
        self.assertEqual(refreshed["items"][0]["id"], selected_id)
        self.assertEqual(self.catalog.search()["total"], 1)
        self.catalog.stage([selected_id], self.root / "stage", ["supply_price"])
        data = json.loads((self.root / "stage/catalog-data.json").read_text(encoding="utf-8"))
        self.assertEqual(data[0]["display_prices"], [{"label": "供货价", "value": 99}])

    def test_duplicate_row_numbers_across_files_remain_distinct(self):
        other = self.root / "another.xlsx"
        workbook(other, {"数码": [["产品名称", "供货价"], ["H1 蓝牙耳机", 90]]})
        self.catalog.index([self.source, other])
        result = self.catalog.search(query="H1")
        ids = [x["id"] for x in result["items"]]
        self.assertEqual(len(set(ids)), 2)
        self.catalog.stage(ids, self.root / "stage")
        data = json.loads((self.root / "stage/catalog-data.json").read_text(encoding="utf-8"))
        self.assertEqual(len(data), 2)

    def test_invalid_update_preserves_previous_index(self):
        self.catalog.index([self.source])
        workbook(self.source, {"bad": [["不认识的列"], ["数据"]]})
        with self.assertRaises(ValueError):
            self.catalog.index([self.source])
        self.assertEqual(self.catalog.stats()["products"], 4)

    def test_config_rules_reindex_and_source_category_wins(self):
        self.catalog.index([self.source])
        self.catalog.close()
        config = {"category_rules": [{"category": "错误分类", "keywords": ["耳机"]}]}
        self.catalog = Catalog(self.root / "index", config)
        self.assertEqual(self.catalog.search()["total"], 0)
        self.assertEqual(self.catalog.index()["updated"], 1)
        self.assertEqual(self.catalog.search(query="耳机")["items"][0]["category"], "数码/耳机")

    def test_hash_recheck_catches_same_stat_tampering(self):
        self.catalog.index([self.source])
        selected = self.catalog.search(query="H1")["items"][0]["id"]
        before = self.source.stat()
        data = bytearray(self.source.read_bytes())
        data[30] ^= 1
        self.source.write_bytes(data)
        os.utime(self.source, ns=(before.st_atime_ns, before.st_mtime_ns))
        self.assertEqual(self.catalog.stale_sources(), [])  # Fast search uses stat only.
        with self.assertRaisesRegex(ValueError, "changed|stale"):
            self.catalog.stage([selected], self.root / "stage")  # Export always hashes.

    def test_image_cache_can_be_repaired_and_three_prices_require_selection(self):
        picture = self.root / "test.png"
        Image.new("RGB", (8, 8), (200, 50, 50)).save(picture)
        book = Workbook()
        sheet = book.active
        sheet.append(["名称", "供货价", "零售价", "批发价", "产品图"])
        sheet.append(["样品", 10, 20, 15])
        sheet.add_image(SheetImage(str(picture)), "E2")
        book.save(self.source)
        self.catalog.index([self.source])
        selected = self.catalog.search()["items"][0]["id"]
        with self.assertRaisesRegex(ValueError, "More than 2"):
            self.catalog.stage([selected], self.root / "stage")
        item = self.catalog.details([selected])[0]
        cached = Path(item["product_image"])
        cached.write_bytes(b"broken")
        with self.assertRaisesRegex(ValueError, "Cached image"):
            self.catalog.stage([selected], self.root / "stage", ["零售价"])
        self.assertEqual(self.catalog.index(rebuild=True)["updated"], 1)
        self.catalog.stage([selected], self.root / "stage", ["零售价"])
        self.assertEqual(cached.read_bytes(), picture.read_bytes())

    def test_per_source_aliases_details_and_uncached_formula(self):
        workbook(self.source, {"数码": [["名称", "供货价", "成本价", "功率"], ["样品", 10, "=5+3", "600W"]]})
        self.catalog.index([self.source])
        self.assertEqual(set(self.catalog.details([self.catalog.search()['items'][0]['id']])[0]['prices']), {'supply_price', 'cost_price'})
        self.catalog.close()
        config = {"source_schemas": [{"match": "supplier.xlsx", "aliases": {"agent_price": ["供货价"]},
                                       "detail_fields": ["功率"]}]}
        self.catalog = Catalog(self.root / "index", config)
        self.catalog.index([self.source])
        selected = self.catalog.search()["items"][0]["id"]
        item = self.catalog.details([selected])[0]
        self.assertEqual(item["features"], "功率：\n600W")
        self.assertIsNone(item["prices"]["cost_price"]["value"])
        self.assertTrue(any("公式无缓存值" in issue for issue in item["issues"]))


if __name__ == "__main__":
    unittest.main()
