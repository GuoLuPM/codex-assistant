import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pool_store import Pool


class PoolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "source.xlsx"
        book = Workbook()
        book.active.append(["名称", "零售价", "卖点", "品类"])
        book.active.append(["红色杯子礼盒", 88, "红色礼盒包装，适合春节赠礼", "杯壶"])
        book.save(self.source)
        self.store = Pool(self.root / "pool")

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_duplicate_bytes_skip_even_after_rename_and_original_removed(self):
        first = self.store.add(self.source)
        copy = self.root / "another-name.xlsx"
        shutil.copyfile(self.source, copy)
        again = self.store.add(copy)
        self.assertEqual(again["file_id"], first["file_id"])
        self.assertEqual(again["status"], "duplicate")
        self.assertEqual(self.store.stats()["products"], 1)
        item = self.store.search()["items"][0]
        self.source.unlink()
        copy.unlink()
        self.store.stage([item["id"]], self.root / "stage", ["零售价"])

    def test_new_version_appends_and_pool_move_preserves_ids(self):
        self.store.add(self.source)
        first = self.store.search()["items"][0]["id"]
        book = Workbook()
        book.active.append(["名称", "零售价"])
        book.active.append(["红色杯子礼盒", 99])
        book.save(self.source)
        self.store.add(self.source)
        self.assertEqual(self.store.stats()["products"], 2)
        self.store.close()
        moved = self.root / "moved pool"
        (self.root / "pool").rename(moved)
        self.store = Pool(moved)
        self.assertEqual(len(self.store.details([first], verify_fresh=True)), 1)
        self.store.stage([first], self.root / "stage", ["零售价"])

    def test_unrecognized_sheet_requires_review_before_file_is_ready(self):
        book = Workbook()
        book.active.title = "产品"
        book.active.append(["名称", "零售价"])
        book.active.append(["合成杯子", 88])
        other = book.create_sheet("待核对")
        other.append(["物件", "报价说明"])
        other.append(["另一款合成商品", "报价需另行核对"])
        book.save(self.source)
        result = self.store.add(self.source)
        self.assertEqual(result["status"], "needs_mapping")
        self.assertEqual(result["skipped_sheets"], ["待核对"])
        self.assertEqual(self.store.document(result["file_id"])["status"], "pending")
        self.assertEqual(self.store.stats()["products"], 0)
        mapping = self.root / "map.json"
        mapping.write_text(json.dumps({"version": 1, "ignore_sheets": ["待核对"], "tables": [{
            "sheet": "产品", "range": "A2:B2", "header_row": 1, "extend_rows": True,
            "fields": {"name": "A"}, "prices": {"retail_price": "B"},
            "expect": {"A1": "名称", "B1": "零售价"}}]}), encoding="utf-8")
        reviewed = self.store.add(self.source, mapping)
        self.assertEqual(reviewed["status"], "ready")
        self.assertEqual(reviewed["products_added"], 1)
        self.assertEqual(self.store.add(self.source)["status"], "duplicate")

    def test_semantic_tags_are_evidenced_and_filter_with_price(self):
        self.store.add(self.source)
        product = self.store.search()["items"][0]["id"]
        tag = {"id": product, "kind": "occasion", "value": "新年", "origin": "inferred",
               "reason": "春节赠礼与新年送礼场景相符", "evidence": "适合春节赠礼"}
        self.store.annotate([tag])
        result = self.store.search(price_field="零售价", maximum=90, constraints={"tags": ["occasion:新年"]})
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["items"][0]["tags"][0]["origin"], "inferred")
        self.store.annotate([{**tag, "kind": "category", "value": "礼品/杯壶"}])
        self.assertEqual(self.store.search(category="礼品")["total"], 1)
        categories = self.store.facets("category")
        self.assertIn("礼品/杯壶", [r["value"] for r in categories["items"]])
        self.assertEqual(self.store.stats()["category_count"], categories["total"])
        with self.assertRaisesRegex(ValueError, "evidence"):
            self.store.annotate([{**tag, "evidence": "虚构信息"}])
        with self.assertRaisesRegex(ValueError, "price.field"):
            self.store.search(maximum=90)

    def test_single_product_import_is_idempotent_and_conflicting_change_fails(self):
        note = self.root / "note.txt"
        note.write_text("便携水杯\n零售价 120元\n重量轻", encoding="utf-8")
        document = self.store.add(note)
        entries = self.store.evidence(document["file_id"])
        refs = {e["text"]: e["id"] for e in entries if e["kind"] == "text"}
        record = {"key": "cup", "fields": {"name": {"ref": refs["便携水杯"]}},
                  "prices": {"retail_price": {"label": {"ref": refs["零售价 120元"], "quote": "零售价"},
                                               "value": {"ref": refs["零售价 120元"], "quote": "120元"}}},
                  "features": [{"ref": refs["重量轻"]}]}
        self.assertEqual(self.store.import_records(document["file_id"], [record])["added"], 1)
        self.assertEqual(self.store.import_records(document["file_id"], [record])["skipped"], 1)
        changed = {**record, "features": []}
        with self.assertRaisesRegex(ValueError, "conflict"):
            self.store.import_records(document["file_id"], [changed])
        bad = {**record, "key": "bad", "fields": {"name": {"ref": refs["便携水杯"], "quote": "杜撰杯"}}}
        with self.assertRaisesRegex(ValueError, "quote"):
            self.store.import_records(document["file_id"], [{**record, "key": "new"}, bad])
        self.assertEqual(self.store.stats()["products"], 1)

    def test_evidence_cache_is_managed_by_store_and_observation_keeps_page(self):
        note = self.root / "note.txt"
        note.write_text("原始商品", encoding="utf-8")
        document = self.store.add(note)
        original = self.store.evidence(document["file_id"])
        cache = self.store.source(document["file_id"]).parent / "evidence.local.json"
        cache.write_text('[{"id":"line:1","kind":"text","text":"伪造文字"}]', encoding="utf-8")
        self.assertEqual(self.store.evidence(document["file_id"]), original)

        import pymupdf
        pdf = self.root / "scan.pdf"
        book = pymupdf.open()
        book.new_page()
        book.new_page()
        book.save(pdf)
        book.close()
        file_id = self.store.add(pdf)["file_id"]
        preview = self.store.preview(file_id, 2)
        self.store.observe(file_id, preview["id"], "逐字观察", "test")
        visual = next(e for e in self.store.evidence(file_id) if e.get("verification"))
        self.assertEqual(visual["page"], 2)


if __name__ == "__main__":
    unittest.main()
