import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pool_store import Pool
from pool_selection import Selections


class LifecycleTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.pool = Pool(self.root / 'pool')
        self.files, self.ids = [], []
        for date, price in [('2026-08-01', 90), ('2026-09-01', 80)]:
            book = Workbook()
            book.active.append(['名称', '型号', '品牌', '零售价', '产品说明'])
            book.active.append(['随行水杯', 'C-500', '远行', price, '容量：500ml\n单位：只\n含税不含运\n起订量：12只'])
            path = self.root / (date + '.xlsx')
            book.save(path)
            fid = self.pool.add(path)['file_id']
            self.files.append(fid)
            self.ids.append(self.pool.search(constraints={'file_id': fid})['items'][0]['id'])

    def tearDown(self):
        self.pool.close()
        self.tmp.cleanup()

    def source(self, index, **changes):
        date = ['2026-08-01', '2026-09-01'][index]
        return {'id': self.files[index], 'revision': 0, 'changes': {'series': 'supplier-cups', 'issued_on': date, **changes},
                'reason': '来源日期及报价系列已核对', 'evidence': [{'ref': 'filename', 'quote': date}]}

    def test_latest_and_explicit_supersession_preserve_records(self):
        self.pool.update_sources([self.source(0), self.source(1)])
        found = self.pool.search(constraints={'latest': True})
        self.assertEqual([p['id'] for p in found['items']], [self.ids[1]])
        update = self.source(1, supersedes={'id': self.files[0], 'revision': 1})
        update['revision'] = 1
        self.pool.update_sources([update])
        self.assertEqual(self.pool.search()['total'], 1)
        self.assertEqual(self.pool.search(constraints={'include_history': True})['total'], 2)
        self.assertEqual(self.pool.db.execute('select count(*) from products').fetchone()[0], 2)
        with self.assertRaisesRegex(ValueError, 'revision'):
            self.pool.update_sources([self.source(0, status='withdrawn')])

    def test_unknown_tied_dates_and_partial_source_do_not_claim_latest(self):
        self.assertEqual(self.pool.search(constraints={'latest': True})['total'], 0)
        with self.assertRaisesRegex(ValueError, 'date evidence'):
            self.pool.update_sources([self.source(0, issued_on='2030-01-01')])
        self.pool.update_sources([self.source(0), self.source(1)])
        with self.pool.db:
            self.pool.db.execute("UPDATE pool_documents SET status='partial' WHERE id=?", (self.files[1],))
        self.assertEqual(self.pool.search(constraints={'latest': True})['total'], 0)
        entry = self.source(1, supersedes={'id': self.files[0], 'revision': 1})
        entry['revision'] = 1
        with self.assertRaisesRegex(ValueError, 'partial'):
            self.pool.update_sources([entry])

    def test_identity_and_terms_are_explicit_and_selected_context_is_checked(self):
        self.pool.link_products([{'group': 'cup-c500', 'ids': self.ids, 'reason': '品牌型号一致，已核对规格',
                                  'evidence': {pid: 'C-500' for pid in self.ids}}])
        self.assertEqual(self.pool.details(self.ids)[0]['product_group'], 'cup-c500')
        self.pool.set_offer_terms([{'id': self.ids[0], 'field': 'retail_price', 'revision': 0,
                                    'terms': {'unit': {'value': '只', 'quote': '单位：只'}}}])
        selections = Selections(self.pool)
        session = selections.create(self.ids[:1], ['零售价'])['session_id']
        selections.select(session, self.ids[:1], 0)
        self.pool.update_sources([self.source(0, status='withdrawn')])
        with self.assertRaisesRegex(ValueError, 'withdrawn|context'):
            selections.seal(session)
        with self.assertRaisesRegex(ValueError, 'quote|evidence'):
            self.pool.set_offer_terms([{'id': self.ids[1], 'field': 'retail_price', 'revision': 0,
                                        'terms': {'unit': {'value': '箱', 'quote': '单位：只'}}}])

    def test_equal_issue_dates_are_ambiguous_and_metadata_changes_block_export(self):
        book = Workbook()
        book.active.append(['名称', '零售价'])
        book.active.append(['另一份报价', 75])
        path = self.root / '修订 2026-09-01.xlsx'
        book.save(path)
        fid = self.pool.add(path)['file_id']
        self.pool.update_sources([self.source(0), self.source(1), {
            'id': fid, 'revision': 0, 'changes': {'series': 'supplier-cups', 'issued_on': '2026-09-01'},
            'reason': '日期相同，需要进一步确认', 'evidence': [{'ref': 'filename', 'quote': '2026-09-01'}]}])
        self.assertEqual(self.pool.search(constraints={'latest': True})['total'], 0)
        self.assertEqual(self.pool.quality()['ambiguous_quote_series'], 1)
        selections = Selections(self.pool)
        session = selections.create(self.ids[:1], ['零售价'])['session_id']
        selections.select(session, self.ids[:1], 0)
        self.pool.derive(self.ids[:1])
        with self.assertRaisesRegex(ValueError, 'context changed'):
            selections.seal(session)
        # Existing installations did not store context hashes. Their original
        # choices stay usable, with source and immutable-fact validation intact.
        with self.pool.db:
            self.pool.db.execute('UPDATE pool_choices SET context_hash=NULL WHERE session_id=?', (session,))
        self.assertEqual(selections.seal(session)['selected_ids'], self.ids[:1])
        entry = self.source(0, status='withdrawn')
        entry['revision'] = 1
        self.pool.update_sources([entry])
        with self.assertRaisesRegex(ValueError, 'withdrawn'):
            selections.seal(session)

    def test_trade_terms_need_evidence_and_do_not_match_unknowns(self):
        terms = {'tax': {'value': 'included', 'quote': '含税不含运'},
                 'shipping': {'value': 'excluded', 'quote': '含税不含运'},
                 'unit': {'value': '只', 'quote': '单位：只'},
                 'moq': {'value': 12, 'quote': '起订量：12只'}}
        self.pool.set_offer_terms([{'id': self.ids[0], 'field': 'retail_price', 'revision': 0, 'terms': terms}])
        result = self.pool.retrieve({'version': 1, 'hard': {'price_field': 'retail_price', 'minimum': 90, 'maximum': 90,
                                   'price_terms': {'unit': '只', 'tax': 'included', 'shipping': 'excluded'}}, 'lanes': [{}]})
        self.assertEqual([p['id'] for p in result['items']], self.ids[:1])
        for fact in [{'currency': {'value': 'CNY', 'quote': '含税不含运'}},
                     {'shipping': {'value': 'included', 'quote': '含税不含运'}},
                     {'moq': {'value': 500, 'quote': '容量：500ml'}}]:
            with self.assertRaisesRegex(ValueError, 'evidence'):
                self.pool.set_offer_terms([{'id': self.ids[1], 'field': 'retail_price', 'revision': 0, 'terms': fact}])

    def test_reviewed_group_offers_can_be_paged_without_hiding_later_quotes(self):
        for price in (70, 60, 50):
            book = Workbook()
            book.active.append(['名称', '型号', '品牌', '零售价'])
            book.active.append(['随行水杯', 'C-500', '远行', price])
            path = self.root / (str(price) + '.xlsx')
            book.save(path)
            fid = self.pool.add(path)['file_id']
            self.ids.append(self.pool.search(constraints={'file_id': fid})['items'][0]['id'])
        self.pool.link_products([{'group': 'cup-c500', 'ids': self.ids, 'reason': '合成同规格商品',
                                  'evidence': {pid: 'C-500' for pid in self.ids}}])
        result = self.pool.retrieve({'version': 1, 'hard': {'price_field': 'retail_price'}, 'lanes': [{}]})
        self.assertEqual(len(result['items']), 1)
        self.assertTrue(result['items'][0]['alternatives_more'])
        self.assertIn('group_offers_clipped_use_paginated_search', result['warnings'])
        paged = []
        for offset in range(0, 5, 2):
            page = self.pool.search(price_field='retail_price', constraints={'product_group': 'cup-c500'}, limit=2, offset=offset)
            paged.extend(p['id'] for p in page['items'])
        self.assertEqual(set(paged), set(self.ids))
        changed = self.pool.derive(self.ids[:1])
        self.assertEqual(changed['updated'], 1)
        self.assertEqual(self.pool.db.execute('SELECT count(*) FROM pool_product_links').fetchone()[0], 0)
        events = self.pool.meta.history('group', 'cup-c500')['items']
        self.assertEqual(events[0]['event_type'], 'identity_changed_requires_review')


if __name__ == '__main__': unittest.main()
