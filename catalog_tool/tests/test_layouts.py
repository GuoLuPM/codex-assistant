import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from extract import read_products
from catalog_store import Catalog
from catalog import view_details
from inspect_source import inspect_source


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def test_explicit_semantics_merged_records_and_price_conflict(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "商品"
        sheet.append(["序号", "品牌", "颜色", "成本价", "零售价"])
        sheet.append([1, "同款产品", "白色", 12, 30])
        sheet.append([None, None, "黑色", None, None])
        for column in "ABDE":
            sheet.merge_cells(f"{column}2:{column}3")
        book.save(self.source)
        mapping = {"version": 1, "tables": [{"sheet": "商品", "range": "A2:E3", "header_row": 1,
                   "fields": {"name": "B", "serial": "A"}, "prices": {"cost_price": "D", "retail_price": "E"},
                   "details": ["C"], "group_by": "name", "extend_rows": True, "expect": {"B1": "品牌", "D1": "成本价"}}]}
        items, _ = read_products(self.source, self.root / "images", {"mapping": mapping})
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["name"], "同款产品")
        self.assertIsNone(items[0]["brand"])
        self.assertIn("白色", items[0]["features"])
        self.assertIn("黑色", items[0]["features"])
        self.assertEqual(items[0]["field_cells"]["name"], ["B2"])
        self.assertEqual(items[0]["field_cells"]["features"], ["C2", "C3"])
        self.assertEqual(items[0]["prices"]["cost_price"]["value"], 12)
        sheet.unmerge_cells("D2:D3")
        sheet["D3"] = 15
        book.save(self.source)
        with self.assertRaisesRegex(ValueError, "Conflicting"):
            read_products(self.source, self.root / "images", {"mapping": mapping})

    def test_repeated_blocks_and_cell_image_provenance(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "卡片"
        for offset, name, price in [(0, "样品甲", 88), (3, "样品乙", 99)]:
            sheet.cell(2, 2 + offset, name)
            sheet.cell(3, 2 + offset, "容量：500ml")
            sheet.cell(4, 2 + offset, "参考价B：")
            sheet.cell(4, 3 + offset, price)
        sheet["A2"] = '=DISPIMG("test-image",1)'
        book.save(self.source)
        blob = io.BytesIO()
        Image.new("RGB", (10, 10), "red").save(blob, format="PNG")
        with ZipFile(self.source, "a", ZIP_DEFLATED) as archive:
            archive.writestr("xl/cellimages.xml", '<root xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><cellImage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="1" name="test-image"/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic></cellImage></root>')
            archive.writestr("xl/_rels/cellimages.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="media/test.png" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"/></Relationships>')
            archive.writestr("xl/media/test.png", blob.getvalue())
        mapping = {"version": 1, "blocks": [{"sheet": "卡片", "range": "A2:F4", "height": 3, "width": 3, "expected_last_row": 4,
                   "fields": {"name": "B1"}, "prices": {"reference_price_b": {"cell": "C3", "label_cell": "B3"}},
                   "details": [{"cell": "B2", "label": "规格"}], "images": {"product_image": ["A1"]},
                   "expect": {"B3": "参考价B："}}]}
        items, _ = read_products(self.source, self.root / "images", {"mapping": mapping})
        self.assertEqual([x["name"] for x in items], ["样品甲", "样品乙"])
        self.assertEqual(items[1]["prices"]["reference_price_b"]["value"], 99)
        self.assertEqual(items[1]["field_cells"]["reference_price_b"], ["F4"])
        self.assertEqual(Path(items[0]["product_image"]).read_bytes(), blob.getvalue())
        self.assertEqual(items[0]["images"][0]["cell"], "A2")

    def test_map_reuse_and_header_drift_are_visible(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["项目", "成交额"])
        sheet.append(["样品", 88])
        book.save(self.source)
        mapping = {"version": 1, "tables": [{"sheet": "Sheet", "range": "A2:B2", "header_row": 1, "extend_rows": True,
                   "fields": {"name": "A"}, "prices": {"quoted_price": "B"}, "expect": {"A1": "项目", "B1": "成交额"}}]}
        map_path = self.root / "source.local.json"
        map_path.write_text(json.dumps(mapping), encoding="utf-8")
        store = Catalog(self.root / "index")
        try:
            store.index([self.source], mapping_path=map_path)
            self.assertEqual(store.index()["skipped"], 1)
            sheet["B2"] = 99
            book.save(self.source)
            store.index()
            self.assertEqual(store.search(price_field="quoted_price")["items"][0]["price"]["value"], 99)
            sheet["B1"] = "另一个口径"
            book.save(self.source)
            with self.assertRaisesRegex(ValueError, "Mapping expectation"):
                store.index()
            self.assertEqual(store.search()["total"], 0)
        finally:
            store.close()

    def test_multiple_images_keep_source_membership(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["产品名称", "产品图片"])
        sheet.append(["样品", None])
        for i, color in enumerate(["red", "blue"]):
            path = self.root / f"{i}.png"
            Image.new("RGB", (10, 10), color).save(path)
            sheet.add_image(SheetImage(str(path)), "B2")
        book.save(self.source)
        items, _ = read_products(self.source, self.root / "images")
        self.assertEqual(len(items[0]["images"]), 2)
        self.assertEqual(items[0]["product_image"], items[0]["images"][0]["path"])

    def test_append_and_multi_region_layout_do_not_drop_or_misprice_rows(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["商品", "卖价", "进价"])
        sheet.append(["样品甲", 20, 12])
        sheet.append(["商品", "进价", "卖价"])
        sheet.append(["样品乙", 13, 22])
        book.save(self.source)
        plans = [
            {"sheet": "Sheet", "range": "A2:C2", "header_row": 1, "expected_last_row": 4,
             "fields": {"name": "A"}, "prices": {"retail_price": "B"}, "expect": {"B1": "卖价"}},
            {"sheet": "Sheet", "range": "A4:C4", "header_row": 3, "extend_rows": True,
             "fields": {"name": "A"}, "prices": {"retail_price": "C"}, "expect": {"C3": "卖价"}},
        ]
        schema = {"mapping": {"version": 1, "tables": plans}}
        items, _ = read_products(self.source, self.root / "images", schema)
        self.assertEqual([p["prices"]["retail_price"]["value"] for p in items], [20, 22])
        sheet.append(["新增样品", 15, 25])
        book.save(self.source)
        with self.assertRaisesRegex(ValueError, "extent changed"):
            read_products(self.source, self.root / "images", schema)
        plans[0]["expected_last_row"] = 5
        items, _ = read_products(self.source, self.root / "images", schema)
        self.assertEqual(len(items), 3)
        self.assertEqual(items[-1]["prices"]["retail_price"]["value"], 25)
        plans[0]["range"] = "A2:C4"
        with self.assertRaisesRegex(ValueError, "Overlapping"):
            read_products(self.source, self.root / "images", schema)

    def test_observation_and_details_are_bounded_without_changing_stored_text(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["产品名称", "功能特点", "零售价(元)", "内部备注"])
        original = "完整说明" * 300
        sheet.append(["样品", original, 100, "不应默认返回的内部内容"])
        book.save(self.source)
        observed = inspect_source(self.source, "Sheet", "A1:D2", max_cells=2, max_chars=20)
        self.assertEqual(len(observed["cells"]), 2)
        self.assertTrue(observed["more"])
        self.assertEqual(inspect_source(self.source)["sheets"][0]["last_value_row"], 2)
        with self.assertRaisesRegex(ValueError, "50 rows"):
            inspect_source(self.source, "Sheet", "A1:D5000")
        store = Catalog(self.root / "index")
        try:
            indexed = store.index([self.source])
            self.assertEqual(indexed["missing_fields"]["prices"], 1)
            self.assertIn("零售价(元)", indexed["unmapped_headers"][0]["fields"])
            selection = store.search()["items"][0]["id"]
            items = store.details([selection])
            short = view_details(items)[0]
            self.assertLessEqual(len(short["features"]), 400)
            self.assertNotIn("raw_fields", short)
            self.assertTrue(short["text_windows"]["features"]["more"])
            second = view_details(items, fields=["features"], max_chars=400, text_offset=400)[0]
            self.assertEqual(second["features"], original[400:800])
            store.stage([selection], self.root / "stage")
            staged = json.loads((self.root / "stage/catalog-data.json").read_text(encoding="utf-8"))
            self.assertEqual(staged[0]["features"], original)
            self.assertEqual(store.sources()["items"][0]["products"], 1)
            self.assertEqual(store.search(source=self.root / "absent.xlsx")["total"], 0)
        finally:
            store.close()

    def test_variants_and_repeated_price_labels_keep_distinct_identities(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["产品", "配置", "价格", "价格"])
        sheet.append(["组合", "整套", 25, 40])
        sheet.append([None, "单独包装", 1, 2])
        sheet.merge_cells("A2:A3")
        book.save(self.source)
        mapping = {"version": 1, "tables": [{"sheet": "Sheet", "range": "A2:D3", "header_row": 1,
                   "extend_rows": True, "fields": {"name": "A", "variant": "B"},
                   "prices": {"cost_price": "C", "retail_price": "D"},
                   "expect": {"A1": "产品", "B1": "配置", "C1": "价格", "D1": "价格"}}]}
        path = self.root / "source.local.json"
        path.write_text(json.dumps(mapping), encoding="utf-8")
        store = Catalog(self.root / "index")
        try:
            store.index([self.source], mapping_path=path)
            item = store.search(query="包装", price_field="retail_price")["items"][0]
            self.assertEqual(item["variant"], "单独包装")
            self.assertEqual(item["price"]["value"], 2)
            with self.assertRaisesRegex(ValueError, "ambiguous"):
                store.stage([item["id"]], self.root / "stage", ["价格"])
            for col in range(2, 5):
                sheet.cell(2, col).value, sheet.cell(3, col).value = sheet.cell(3, col).value, sheet.cell(2, col).value
            book.save(self.source)
            store.index()
            self.assertEqual(store.search(query="包装")["items"][0]["id"], item["id"])
        finally:
            store.close()

    def test_excel_price_errors_cannot_be_published_as_quotes(self):
        book = Workbook()
        book.active.append(["产品名称", "供货价"])
        book.active.append(["样品", "#REF!"])
        book.save(self.source)
        store = Catalog(self.root / "index")
        try:
            store.index([self.source])
            product = store.search()["items"][0]
            self.assertEqual(product["prices"]["supply_price"]["error"], "spreadsheet_error")
            with self.assertRaisesRegex(ValueError, "spreadsheet error"):
                store.stage([product["id"]], self.root / "stage")
        finally:
            store.close()

    def test_mapping_typos_fail_instead_of_silently_dropping_fields(self):
        book = Workbook()
        book.active.append(["品名", "图片"])
        book.active.append(["样品", None])
        book.save(self.source)
        plan = {"sheet": "Sheet", "range": "A2:B2", "header_row": 1, "extend_rows": True,
                "fields": {"name": "A"}, "expect": {"A1": "品名"}, "image": {"product_image": ["B"]}}
        schema = {"mapping": {"version": 1, "tables": [plan]}}
        with self.assertRaisesRegex(ValueError, "Unknown layout"):
            read_products(self.source, self.root / "images", schema)
        plan["images"] = plan.pop("image")
        plan["images"]["product_image"] = "B"
        with self.assertRaisesRegex(ValueError, "lists of references"):
            read_products(self.source, self.root / "images", schema)


if __name__ == "__main__":
    unittest.main()
