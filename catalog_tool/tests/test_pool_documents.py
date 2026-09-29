import io
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import pymupdf as fitz

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pool_documents import read_evidence, inspect_evidence, render_page


class EvidenceTests(unittest.TestCase):
    def test_native_formats_images_and_bounded_refs(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            assets = root / "assets"
            text = root / "source.txt"
            text.write_text("名称\n" + "原文" * 300, encoding="utf-8")
            entries = read_evidence(text, assets)
            self.assertEqual(entries[0]["text"], "名称")
            bounded = inspect_evidence(entries, limit=1, offset=1, max_chars=30)
            self.assertTrue(bounded["items"][0]["truncated"])
            self.assertEqual(len(bounded["items"][0]["text"]), 30)
            book = Workbook()
            book.active.append(["名称", "价格"])
            book.active.append(["样品", 88])
            xlsx = root / "source.xlsx"
            book.save(xlsx)
            self.assertEqual(read_evidence(xlsx, assets)[2]["locator"], "Sheet!A2")
            pdf = root / "source.pdf"
            doc = fitz.open()
            doc.new_page().insert_text((30, 30), "Retail 88")
            doc.save(pdf)
            doc.close()
            self.assertIn("Retail 88", read_evidence(pdf, assets)[0]["text"])
            self.assertTrue(Path(render_page(pdf, assets, 1)["path"]).is_file())
            png = io.BytesIO()
            Image.new("RGB", (20, 30), "red").save(png, format="PNG")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Source product\nRetail 99"
            slide.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(2))
            ppt = root / "source.pptx"
            deck.save(ppt)
            evidence = read_evidence(ppt, assets)
            self.assertEqual(evidence[0]["text"], "Source product\nRetail 99")
            self.assertEqual(Path(evidence[1]["path"]).read_bytes(), png.getvalue())
            self.assertEqual(evidence, read_evidence(ppt, assets))


if __name__ == "__main__":
    unittest.main()
