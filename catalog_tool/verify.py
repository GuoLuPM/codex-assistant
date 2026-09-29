"""Compare generated slide text with extracted workbook values."""

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from catalog_store import check_sources


def compact(value):
    return re.sub(r"\s+", "", str(value))


def verify(work_dir: Path, output: Path):
    data = json.loads((work_dir / "catalog-data.json").read_text(encoding="utf-8"))
    products = {item["id"]: item for item in data}
    assert len(products) == len(data), "Duplicate product IDs"
    manifest = json.loads((work_dir / "catalog-manifest.json").read_text(encoding="utf-8"))
    deck = Presentation(output)
    assert len(deck.slides) == len(manifest), "Slide count differs from manifest"
    parts = defaultdict(list)
    for meta, slide in zip(manifest, deck.slides):
        item = products[meta["product_id"]]
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text]
        title = re.sub(r"\s+", " ", item["name"]).strip()
        if item["model"] is not None and str(item["model"]).strip() and str(item["model"]) not in title:
            title += " / " + re.sub(r"\s+", " ", str(item["model"])).strip()
        assert texts[0] == title, (meta, "product name")
        assert f"序号 {item['serial'] if item['serial'] is not None else '未提供'}" in texts, (meta, "serial")
        prices = item["display_prices"]
        expected = "\n".join(f"{p['label']}  {'未提供' if p['value'] is None or not str(p['value']).strip() else p['value']}" for p in prices)
        assert texts[-1] == expected, (meta, "price")
        notes = slide.notes_slide.notes_text_frame.text
        assert item["source_file"] in notes and item["id"] in notes, (meta, "provenance")
        parts[item["id"]].append(texts[-2])

    for row, item in products.items():
        original = "表格未提供功能特点" if not item["features"] or not str(item["features"]).strip() else str(item["features"]).strip()
        assert compact("".join(parts[row])) == compact(original), (row, "features")
    check_sources(data)
    return len(products), len(deck.slides)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    products, slides = verify(args.work_dir.resolve(), args.output.resolve())
    print(f"Verified {products} products across {slides} slides")
