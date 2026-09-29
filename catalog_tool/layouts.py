"""将显式行表/重复块映射转成统一的来源单元格记录，不判断商品语义。"""

import re
from dataclasses import dataclass

from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple, get_column_letter, column_index_from_string


def empty(value):
    return value is None or isinstance(value, str) and not value.strip()


class SheetValues:
    def __init__(self, sheet):
        self.sheet = sheet
        self.merges = list(sheet.merged_cells.ranges)

    def resolve(self, row, col):
        for merged in self.merges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                row, col = merged.min_row, merged.min_col
                break
        return self.sheet.cell(row, col)

    def span(self, row, col):
        return next((m for m in self.merges if m.min_row <= row <= m.max_row and m.min_col <= col <= m.max_col), None)


@dataclass
class RecordLayout:
    values: SheetValues
    plan: dict
    top: int
    bottom: int
    left: int
    right: int
    block: bool = False

    def cells(self, ref):
        if isinstance(ref, dict):
            ref = ref.get("cell") if self.block else ref.get("column")
        if not isinstance(ref, str):
            raise ValueError("Field reference must identify a column or a relative block cell")
        if self.block:
            if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*", ref):
                raise ValueError(f"Invalid relative cell: {ref}")
            row, col = coordinate_to_tuple(ref)
            row, col = self.top + row - 1, self.left + col - 1
            if row > self.bottom or col > self.right:
                raise ValueError(f"Reference {ref} falls outside its block")
            coords = [(row, col)]
        else:
            col = column_index_from_string(ref)
            if not self.left <= col <= self.right:
                raise ValueError(f"Column {ref} falls outside its table range")
            coords = [(row, col) for row in range(self.top, self.bottom + 1)]
        result = {}
        for row, col in coords:
            cell = self.values.resolve(row, col)
            result[cell.coordinate] = cell
        return list(result.values())

    def label(self, ref):
        if isinstance(ref, dict) and "label" in ref:
            return str(ref["label"])
        if isinstance(ref, dict) and "label_cell" in ref:
            cells = self.cells(ref["label_cell"])
            return str(cells[0].value or ref["label_cell"]).strip()
        if self.block:
            return ref if isinstance(ref, str) else ref["cell"]
        col = ref if isinstance(ref, str) else ref["column"]
        header = self.values.resolve(self.plan["header_row"], column_index_from_string(col)).value
        return str(header or col).strip()

    def image_cells(self, ref):
        # 浮动图可能锚定在合并区的次行；保留原锚点及实际合并起点。
        refs = {cell.coordinate for cell in self.cells(ref)}
        if not self.block:
            col = ref if isinstance(ref, str) else ref["column"]
            refs.update(f"{col}{row}" for row in range(self.top, self.bottom + 1))
        return sorted(refs, key=coordinate_to_tuple)


def check_expectations(values, expected, record=None):
    if not expected:
        raise ValueError("Explicit mappings require source-cell expectations")
    for ref, expected_value in expected.items():
        cell = record.cells(ref)[0] if record else values.sheet[ref]
        if cell.value != expected_value:
            raise ValueError(f"Mapping expectation changed at {values.sheet.title}!{cell.coordinate}; inspect and remap")


def plan_bounds(sheet, plan, block=False):
    left, top, right, bottom = range_boundaries(plan["range"])
    if not all((left, top, right, bottom)) or top > bottom or left > right:
        raise ValueError("Mapping needs a bounded rectangular range")
    last_row = max((cell.row for cell in sheet._cells.values() if not empty(cell.value)), default=0)
    if plan.get("extend_rows"):
        if block:
            raise ValueError("Repeated blocks require a fixed range and expected_last_row")
        bottom = max(bottom, max((cell.row for cell in sheet._cells.values()
                                 if left <= cell.column <= right and not empty(cell.value)), default=bottom))
    elif not plan.get("automatic"):
        if plan.get("expected_last_row") != last_row:
            raise ValueError(f"Sheet extent changed or expected_last_row missing in {sheet.title}; inspect and remap")
    return left, top, right, bottom


def records_for_plan(sheet, plan, block=False):
    common = {"sheet", "range", "fields", "prices", "details", "images", "expect", "category_from_sheet",
              "expected_last_row", "automatic"}
    options = {"height", "width"} if block else {"header_row", "extend_rows", "group_by", "exclude_rows", "skip_repeated_headers"}
    if set(plan) - common - options:
        raise ValueError(f"Unknown layout options: {sorted(set(plan) - common - options)}")
    for key in ("fields", "prices", "images", "expect"):
        if not isinstance(plan.get(key, {}), dict):
            raise ValueError(f"{key} must be an object")
    if not isinstance(plan.get("details", []), list) or any(not isinstance(refs, list) for refs in plan.get("images", {}).values()):
        raise ValueError("details and each image role must be lists of references")
    for key in ("extend_rows", "category_from_sheet", "skip_repeated_headers"):
        if key in plan and not isinstance(plan[key], bool):
            raise ValueError(f"{key} must be a boolean")
    if set(plan.get("prices", {})) & {"name", "serial", "model", "variant", "features", "category", "brand", "supplier"}:
        raise ValueError("Price roles must be distinct from product fields")
    values = SheetValues(sheet)
    left, top, right, bottom = plan_bounds(sheet, plan, block)
    if "name" not in plan.get("fields", {}):
        raise ValueError("Every mapped layout requires a name field")
    if set(plan["fields"]) - {"name", "serial", "model", "variant", "category", "brand", "supplier"}:
        raise ValueError("Unsupported mapped scalar field")
    if plan.get("category_from_sheet") and "category" in plan["fields"]:
        raise ValueError("Choose either a source category column or category_from_sheet")
    if block:
        height, width = plan["height"], plan["width"]
        if height <= 0 or width <= 0 or (bottom - top + 1) % height or (right - left + 1) % width:
            raise ValueError("Block range must be divisible by height and width")
        for row in range(top, bottom + 1, height):
            for col in range(left, right + 1, width):
                record = RecordLayout(values, plan, row, row + height - 1, col, col + width - 1, True)
                # 空白尾块可省略，出现数据的块仍严格校验标签。
                refs = [*plan["fields"].values(), *plan.get("prices", {}).values(), *plan.get("details", [])]
                if all(empty(cell.value) for ref in refs for cell in record.cells(ref)):
                    continue
                check_expectations(values, plan.get("expect"), record)
                yield record
        return
    if not isinstance(plan.get("header_row"), int):
        raise ValueError("Table mappings require header_row")
    if not plan.get("automatic"):
        check_expectations(values, plan.get("expect"))
    row = top
    name_col = column_index_from_string(plan["fields"]["name"] if isinstance(plan["fields"]["name"], str) else plan["fields"]["name"]["column"])
    while row <= bottom:
        end = row
        if plan.get("group_by") == "name":
            merged = values.span(row, name_col)
            if merged:
                if merged.min_row != row or merged.max_row > bottom:
                    raise ValueError("Mapped record cuts a merged name region")
                end = merged.max_row
        elif plan.get("group_by") not in (None, "row"):
            raise ValueError("group_by must be row or name")
        if row not in plan.get("exclude_rows", []):
            expected = plan.get("expect", {})
            repeat = plan.get("skip_repeated_headers", False) and expected and all(
                values.sheet.cell(row, coordinate_to_tuple(ref)[1]).value == value for ref, value in expected.items())
            if not repeat:
                yield RecordLayout(values, plan, row, end, left, right)
        row = end + 1
