"""读取 XLSX 包内图片与公式；只使用包内关系，不访问外部地址。"""

import hashlib
import posixpath
import re
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DRAW_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
ART_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
CELL_IMAGE = re.compile(r'^=(?:_xlfn\.)?DISPIMG\("([^"\r\n]+)"\s*,\s*\d+\s*\)$', re.I)


def part_target(owner, target):
    result = posixpath.normpath(posixpath.join(posixpath.dirname(owner), target)) if not target.startswith("/") else target.lstrip("/")
    if result.startswith(("../", "/")) or ":" in result or ".." in result.split("/"):
        raise ValueError("External/unsafe XLSX part reference")
    return result


class WorkbookPackage:
    def __init__(self, source):
        self.archive = ZipFile(source)
        workbook = ET.fromstring(self.archive.read("xl/workbook.xml"))
        links = self.relationships("xl/workbook.xml", "worksheet")
        self.sheets = {node.attrib["name"]: links[node.attrib[f"{{{REL_NS}}}id"]]
                       for node in workbook.findall(f"{{{SHEET_NS}}}sheets/{{{SHEET_NS}}}sheet")}
        self.cell_images = {}
        self._strings = None
        if "xl/cellimages.xml" in self.archive.namelist():
            root = ET.fromstring(self.archive.read("xl/cellimages.xml"))
            links = self.relationships("xl/cellimages.xml", "image")
            for picture in root.findall(f".//{{{DRAW_NS}}}pic"):
                identity = picture.find(f".//{{{DRAW_NS}}}cNvPr")
                blip = picture.find(f".//{{{ART_NS}}}blip")
                if identity is not None and blip is not None:
                    rel = blip.attrib.get(f"{{{REL_NS}}}embed")
                    if rel in links:
                        self.cell_images[identity.attrib["name"]] = links[rel]

    def relationships(self, owner, kind):
        rel_path = posixpath.join(posixpath.dirname(owner), "_rels", posixpath.basename(owner) + ".rels")
        if rel_path not in self.archive.namelist():
            return {}
        root = ET.fromstring(self.archive.read(rel_path))
        return {node.attrib["Id"]: part_target(owner, node.attrib["Target"])
                for node in root if node.attrib.get("TargetMode") != "External" and node.attrib.get("Type", "").endswith("/" + kind)}

    def sheet_info(self, name):
        root = ET.fromstring(self.archive.read(self.sheets[name]))
        formulas = {cell.attrib["r"]: "=" + (formula.text or "")
                    for cell in root.findall(f".//{{{SHEET_NS}}}sheetData/{{{SHEET_NS}}}row/{{{SHEET_NS}}}c")
                    if (formula := cell.find(f"{{{SHEET_NS}}}f")) is not None}
        merges = [node.attrib["ref"] for node in root.findall(f"{{{SHEET_NS}}}mergeCells/{{{SHEET_NS}}}mergeCell")]
        return formulas, merges

    def sheet_extent(self, name):
        """Cached nonempty values, excluding formatting-only cells and empty strings."""
        if self._strings is None:
            self._strings = []
            if "xl/sharedStrings.xml" in self.archive.namelist():
                strings = ET.fromstring(self.archive.read("xl/sharedStrings.xml"))
                self._strings = ["".join(node.itertext()) for node in strings]
        root = ET.fromstring(self.archive.read(self.sheets[name]))
        rows, cols = [], []
        for cell in root.findall(f".//{{{SHEET_NS}}}sheetData/{{{SHEET_NS}}}row/{{{SHEET_NS}}}c"):
            value = cell.findtext(f"{{{SHEET_NS}}}v")
            if cell.attrib.get("t") == "s" and value is not None:
                value = self._strings[int(value)]
            elif cell.attrib.get("t") == "inlineStr":
                value = "".join(t.text or "" for t in cell.findall(f".//{{{SHEET_NS}}}t"))
            if value is not None and value.strip():
                row, col = coordinate_to_tuple(cell.attrib["r"])
                rows.append(row)
                cols.append(col)
        return {"last_value_row": max(rows, default=0),
                "value_range": f"{get_column_letter(min(cols))}{min(rows)}:{get_column_letter(max(cols))}{max(rows)}" if rows else None}

    def close(self):
        self.archive.close()


class ImageReader:
    def __init__(self, sheet, package, image_dir, formulas):
        self.sheet, self.package = sheet, package
        self.image_dir, self.formulas = Path(image_dir), formulas
        self.floating = defaultdict(list)
        self.cache = {}
        self.image_dir.mkdir(parents=True, exist_ok=True)
        for image in sheet._images:
            if hasattr(image.anchor, "_from"):
                anchor = image.anchor._from
                cell = sheet.cell(anchor.row + 1, anchor.col + 1).coordinate
                self.floating[cell].append(image)

    def _save(self, data, suffix, cell, kind):
        suffix = suffix.lower().lstrip(".")
        suffix = "jpg" if suffix == "jpeg" else suffix
        if suffix not in {"png", "jpg"}:
            raise ValueError(f"Unsupported source image format {suffix} at {self.sheet.title}!{cell}")
        digest = hashlib.sha256(data).hexdigest()
        target = self.image_dir / f"{digest}.{suffix}"
        if not target.exists() or hashlib.sha256(target.read_bytes()).hexdigest() != digest:
            target.write_bytes(data)
        return {"path": str(target.resolve()), "sha256": digest, "cell": cell, "kind": kind}

    def read(self, cell):
        if cell in self.cache:
            return self.cache[cell]
        result, issues = [], []
        for image in self.floating[cell]:
            result.append(self._save(image._data(), image.format, cell, "floating"))
        value = self.formulas.get(cell) or self.sheet[cell].value
        match = CELL_IMAGE.fullmatch(str(value or "").strip())
        if match:
            target = self.package.cell_images.get(match[1])
            if target:
                result.append(self._save(self.package.archive.read(target), Path(target).suffix, cell, "cell_image"))
            else:
                issues.append(f"单元格图片未找到:{cell}")
        elif isinstance(value, str) and value.strip() and not result:
            issues.append(f"图片字段文本未解析:{cell}")
        self.cache[cell] = result, issues
        return result, issues
