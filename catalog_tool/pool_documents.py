"""Native document evidence, with stable locators and bounded model-facing reads."""

import csv
import hashlib
import io
from pathlib import Path

from catalog_store import sha


def image_entry(data, assets, ref, locator, page):
    from PIL import Image
    with Image.open(io.BytesIO(data)) as im:
        if im.format not in {"PNG", "JPEG"}:
            out = io.BytesIO()
            im.convert("RGB").save(out, format="PNG")
            data = out.getvalue()
            suffix = ".png"
        else:
            suffix = ".png" if im.format == "PNG" else ".jpg"
    checksum = hashlib.sha256(data).hexdigest()
    assets = Path(assets)
    assets.mkdir(parents=True, exist_ok=True)
    target = assets / (checksum + suffix)
    if not target.exists():
        target.write_bytes(data)
    elif sha(target) != checksum:
        raise ValueError("Evidence image cache hash mismatch")
    return {"id": ref, "kind": "image", "path": str(target.resolve()), "sha256": checksum, "locator": locator, "page": page}


def read_evidence(source, assets, allow_empty=False):
    source = Path(source)
    before = sha(source)
    entries = []

    def text(value, ref, locator, page=1):
        if value is not None and str(value).strip():
            entries.append({"id": ref, "kind": "text", "text": str(value), "locator": locator, "page": page})

    suffix = source.suffix.lower()
    if suffix in {".txt", ".md"}:
        for n, line in enumerate(source.read_text(encoding="utf-8-sig").splitlines(), 1):
            text(line, f"line:{n}", f"line {n}")
    elif suffix == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as stream:
            for r, row in enumerate(csv.reader(stream), 1):
                for c, value in enumerate(row, 1):
                    text(value, f"r{r}c{c}", f"row {r}, column {c}")
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        entries.append(image_entry(source.read_bytes(), assets, "image:1", "original image", 1))
    elif suffix == ".pdf":
        import pymupdf as fitz
        with fitz.open(source) as book:
            for n, page in enumerate(book, 1):
                for b, block in enumerate(page.get_text("blocks"), 1):
                    text(block[4], f"p{n}:text{b}", f"page {n}, block {b}", n)
                for j, image in enumerate(page.get_images(full=True), 1):
                    extracted = book.extract_image(image[0])
                    if not extracted or not extracted.get("image"):
                        raise ValueError(f"Cannot extract PDF image on page {n}")
                    entries.append(image_entry(extracted["image"], assets, f"p{n}:image{j}", f"page {n}, image {j}", n))
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        def shapes(collection, n, prefix=""):
            for j, shape in enumerate(collection, 1):
                ref = f"s{n}:{prefix}shape{j}"
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, n, prefix + f"group{j}:")
                if shape.has_text_frame:
                    text(shape.text, ref, f"slide {n}, {prefix}shape {j}", n)
                if shape.has_table:
                    for r, row in enumerate(shape.table.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            text(cell.text, f"{ref}:r{r}c{c}", f"slide {n}, {prefix}table {j}, row {r}, column {c}", n)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    entries.append(image_entry(shape.image.blob, assets, ref + ":image", f"slide {n}, {prefix}picture {j}", n))
        for n, slide in enumerate(Presentation(source).slides, 1):
            shapes(slide.shapes, n)
    elif suffix == ".docx":
        from docx import Document
        document = Document(source)
        for n, paragraph in enumerate(document.paragraphs, 1):
            text(paragraph.text, f"para:{n}", f"paragraph {n}")
        for t, table in enumerate(document.tables, 1):
            for r, row in enumerate(table.rows, 1):
                for c, cell in enumerate(row.cells, 1):
                    text(cell.text, f"table{t}:r{r}c{c}", f"table {t}, row {r}, column {c}")
    elif suffix in {".xlsx", ".xls"}:
        if suffix == ".xlsx":
            import openpyxl
            book = openpyxl.load_workbook(source, read_only=True, data_only=True)
            try:
                for n, sheet in enumerate(book, 1):
                    for row in sheet:
                        for cell in row:
                            if cell.value is not None:
                                text(cell.value, f"sheet{n}:{cell.coordinate}", f"{sheet.title}!{cell.coordinate}", n)
            finally:
                book.close()
        else:
            import xlrd
            book = xlrd.open_workbook(source)
            try:
                for n, sheet in enumerate(book.sheets(), 1):
                    for r in range(sheet.nrows):
                        for c in range(sheet.ncols):
                            text(sheet.cell_value(r, c), f"sheet{n}:r{r+1}c{c+1}", f"{sheet.name}!R{r+1}C{c+1}", n)
            finally:
                book.release_resources()
    else:
        raise ValueError(f"Unsupported evidence format {suffix}; convert explicitly, preserving the original")
    if sha(source) != before:
        raise ValueError("Source changed during evidence extraction")
    if not entries and not allow_empty:
        raise ValueError("No native evidence extracted; PDF may require render then visual observation")
    return entries


def inspect_evidence(entries, page=None, offset=0, limit=20, max_chars=240):
    if not 1 <= limit <= 50 or offset < 0 or not 20 <= max_chars <= 2000 or (page is not None and page < 1):
        raise ValueError("Invalid evidence pagination")
    entries = [e for e in entries if page is None or e["page"] == page]
    output = []
    for entry in entries[offset:offset+limit]:
        item = dict(entry)
        if len(item.get("text", "")) > max_chars:
            item["text_length"] = len(item["text"])
            item["text"] = item["text"][:max_chars]
            item["truncated"] = True
        output.append(item)
    return {"total": len(entries), "items": output, "offset": offset, "limit": limit}


def render_page(source, assets, page):
    import pymupdf as fitz
    source = Path(source)
    before = sha(source)
    with fitz.open(source) as book:
        if not 1 <= page <= len(book):
            raise ValueError("Page outside document")
        data = book[page - 1].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes("png")
    if sha(source) != before:
        raise ValueError("Source changed while rendering")
    return image_entry(data, assets, f"p{page}:preview", f"page {page} preview (whole page)", page)
