"""Exercise a synthetic product pool; optionally validate a real PPT export."""
import argparse
import importlib
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'catalog_tool'))


def check(ppt=False, require_pdf=False, workspace=False):
    modules = ['openpyxl', 'PIL', 'pptx', 'docx', 'xlrd'] + (['pymupdf'] if require_pdf else []) + (['fastapi', 'uvicorn'] if workspace else [])
    for name in modules: importlib.import_module(name)
    with sqlite3.connect(':memory:') as database: database.execute('CREATE VIRTUAL TABLE probe USING fts5(text)')
    from openpyxl import Workbook
    from pool_store import Pool
    from pool_selection import Selections
    private = ROOT / 'data/deployment-checks'
    private.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix='run-', dir=private))
    result = {'python': sys.executable, 'imports': modules, 'sqlite_fts5': True, 'ppt_verified': False}
    try:
        source = work / 'synthetic.xlsx'
        book = Workbook()
        book.active.append(['名称', '零售价', '功能特点'])
        book.active.append(['部署测试保温杯', 88, '合成测试资料\n容量：500ml\n礼盒包装'])
        book.save(source)
        pool = Pool(work / 'pool')
        try:
            pool.add(source)
            renamed = work / 'renamed.xlsx'
            shutil.copyfile(source, renamed)
            if pool.add(renamed)['status'] != 'duplicate': raise ValueError('Duplicate-file check failed')
            items = pool.search(query='保温杯', price_field='retail_price', maximum=100)['items']
            if len(items) != 1 or items[0]['price']['value'] != 88: raise ValueError('Source price/search mismatch')
            choices = Selections(pool)
            session = choices.create([items[0]['id']], ['retail_price'])['session_id']
            choices.select(session, [items[0]['id']], 0)  # synthetic test only
            if choices.state(session)['selected_ids'] != [items[0]['id']]: raise ValueError('Selection persistence failed')
            if require_pdf:
                import pymupdf
                with pymupdf.open() as pdf:
                    page = pdf.new_page()
                    page.insert_text((40, 40), 'Synthetic PDF product: test cup, price 88')
                    pdf.save(work / 'synthetic.pdf')
                document = pool.add(work / 'synthetic.pdf')
                if not any('price 88' in row.get('text', '') for row in pool.evidence(document['file_id'])):
                    raise ValueError('PDF source extraction failed')
        finally: pool.close()
        if ppt:
            output = work / 'smoke.pptx'
            command = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ROOT / 'assistant.ps1'),
                       'run', 'pool', '--index-dir', str(work / 'pool'), 'ppt', '--session', session, '--output', str(output)]
            done = subprocess.run(command, cwd=ROOT, capture_output=True, env={**os.environ, 'PYTHONIOENCODING': 'utf-8'}, timeout=180)
            if done.returncode: raise ValueError('PPT pipeline failed: ' + (done.stderr or done.stdout).decode('utf-8', errors='replace')[-1500:])
            from pptx import Presentation
            deck = Presentation(output)
            text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            if len(deck.slides) != 1 or '部署测试保温杯' not in text or '零售价  88' not in text: raise ValueError('Exported PPT facts differ')
            result.update(ppt_verified=True, slides=len(deck.slides))
        result.update(pool_import=True, duplicate_skip=True, budget_search=True, selection=True)
        if workspace:
            sys.path.insert(0, str(ROOT))
            from workspace_tool.cli import launch, control, process_alive
            import time
            if not (ROOT / 'workspace_tool/web/dist/index.html').is_file(): raise ValueError('Built workspace page is missing; use the Windows release asset')
            directory = work / 'workspace'
            launch(directory, work / 'pool')
            state = json.loads((directory / 'server.local.json').read_text(encoding='utf-8'))
            try:
                deadline = time.monotonic() + 45
                while time.monotonic() < deadline:
                    ready = control(state, 'status')
                    if ready['ready'] or ready.get('error'): break
                    time.sleep(.5)
                if not ready['ready']: raise ValueError('Codex connection not ready: ' + str(ready.get('error') or 'Check Codex login and installed models'))
                result.update(workspace_ready=True, codex_version=ready['version'], models=ready['models'])
            finally:
                control(state, 'stop', method='POST')
                for _ in range(100):
                    # The registry is removed just before Python exits. On
                    # Windows its log remains locked until the process is gone.
                    if not (directory / 'server.local.json').exists() and not process_alive(state['pid']): break
                    time.sleep(.1)
                else: raise ValueError('Owned workspace has not finished shutting down')
        # Only disposable synthetic state in this exact task directory is removed.
        if not work.resolve().is_relative_to(private.resolve()) or any(p.is_symlink() or (hasattr(p, 'is_junction') and p.is_junction()) for p in work.rglob('*')):
            raise ValueError('Unsafe test cleanup path')
        shutil.rmtree(work)
        return result
    except Exception:
        print(json.dumps({'failed_check_directory': str(work)}, ensure_ascii=False), file=sys.stderr)
        raise


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ppt', action='store_true')
    parser.add_argument('--require-pdf', action='store_true')
    parser.add_argument('--workspace', action='store_true')
    args = parser.parse_args()
    print(json.dumps(check(args.ppt, args.require_pdf, args.workspace), ensure_ascii=False))
